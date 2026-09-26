"""Lifecycle tab: row shaping, per-case flags, and the two routes (no network)."""

import asyncio
import io
import json

import pytest
from fastapi.testclient import TestClient

from creatio_case_lookup import lifecycle, lifecycle_export, server

G = "9027b5a8-ae65-46b6-ba25-c484e7568187"
C = "11111111-2222-3333-4444-555555555555"


def row(status, owner, start, end, group="Team"):
    return {"Status": {"Name": status}, "Owner": {"Name": owner}, "Group": {"Name": group},
            "StartDate": start, "EndDate": end}


RAW = [
    # Out of order on purpose: shape_rows sorts by StartDate.
    row("In progress", "Bo", "2026-01-01T12:00:00Z", "2026-01-02T00:00:00Z"),
    row("New", "Al", "2026-01-01T00:00:00Z", "2026-01-01T10:00:00Z"),
    row("New", "Bo", "2026-01-01T10:00:00Z", "2026-01-01T12:00:00Z"),  # reassignment, same status
    row("Waiting for response", "Bo", "2026-01-02T00:00:00Z", "2026-01-03T00:00:00Z"),
    row("In progress", "Bo", "2026-01-03T00:00:00Z", "2026-01-03T01:00:00Z"),
    row("Waiting for response", "Bo", "2026-01-03T01:00:00Z", "2026-01-03T02:00:00Z"),
    row("In progress", "Bo", "2026-01-03T02:00:00Z", "2026-01-03T03:00:00Z"),
    row("Resolved", "Bo", "2026-01-03T03:00:00Z", "2026-01-04T00:00:00Z"),
    row("Closed", "Bo", "2026-01-04T00:00:00Z", "0001-01-01T00:00:00Z"),  # Creatio's open sentinel
]


def test_shape_rows_sorts_and_normalises_sentinel():
    rows = lifecycle.shape_rows(RAW)
    assert [r["status"] for r in rows][:3] == ["New", "New", "In progress"]
    assert rows[-1]["end"] is None and rows[-1]["minutes"] is None
    assert rows[0]["minutes"] == 600


def test_analyze_case_flags():
    a = lifecycle.analyze_case(lifecycle.shape_rows(RAW))
    assert a["arc"] == ["New", "In progress", "Waiting for response", "In progress",
                        "Waiting for response", "In progress", "Resolved", "Closed"]
    assert a["bounces"] == 2
    assert a["reassignments"] == 1
    assert a["reassignOnlyRows"] == 1
    assert a["fullArc"] is True
    assert a["minutesByStatus"]["Waiting for response"] == 24 * 60 + 60
    assert a["minutesByStatus"]["In progress"] == 12 * 60 + 60 + 60


def test_reassignment_through_unowned_row_counts_once():
    rows = lifecycle.shape_rows([
        row("In progress", "Amy", "2026-01-01T00:00:00Z", "2026-01-01T01:00:00Z"),
        row("In progress", "", "2026-01-01T01:00:00Z", "2026-01-01T02:00:00Z"),
        row("In progress", "Mel", "2026-01-01T02:00:00Z", "0001-01-01T00:00:00Z"),
    ])
    assert lifecycle.analyze_case(rows)["reassignments"] == 1


def test_sample_rejects_entity_outside_allowlist(monkeypatch):
    monkeypatch.setattr(lifecycle, "ALLOWED_ENTITIES", ["Case"])
    with pytest.raises(lifecycle.LifecycleError, match="allowlist"):
        asyncio.run(lifecycle.lifecycle_sample(G, "2025-09-01"))


@pytest.mark.parametrize("gid,since", [("nope", "2025-09-01"), (G, "01/09/2025"), (G, "2025-09-01' or 1 eq 1")])
def test_sample_validates_input(monkeypatch, gid, since):
    monkeypatch.setattr(lifecycle, "ALLOWED_ENTITIES", [])
    with pytest.raises(lifecycle.LifecycleError):
        asyncio.run(lifecycle.lifecycle_sample(gid, since))


def test_sample_pages_cases_and_reads_rows(monkeypatch):
    monkeypatch.setattr(lifecycle, "ALLOWED_ENTITIES", [])
    monkeypatch.setattr(lifecycle, "MAX_TOP", 2)
    calls = []

    async def fake_query(entity, **kw):
        calls.append((entity, kw["filter"]))
        if entity == "Case":
            if "CreatedOn lt" not in kw["filter"]:
                return [{"Id": C, "Number": "SR1", "CreatedOn": "2026-02-02T00:00:00Z", "Owner": {"Name": "Bo"}},
                        {"Id": C, "Number": "SR2", "CreatedOn": "2026-02-01T00:00:00Z"}]
            return [{"Id": C, "Number": "SR3", "CreatedOn": "2026-01-01T00:00:00Z"}]
        return RAW[:2]

    monkeypatch.setattr(lifecycle, "query_records", fake_query)
    out = asyncio.run(lifecycle.lifecycle_sample(G, "2025-09-01", 10))
    assert [c["Number"] for c in out["cases"]] == ["SR1", "SR2", "SR3"]
    assert "CreatedOn lt 2026-02-01T00:00:00Z" in calls[1][1]
    assert all(f == f"Case/Id eq {C}" for e, f in calls if e == "CaseLifecycle")
    assert out["summary"]["cases"] == 3 and out["summary"]["rows"] == 6
    assert out["cases"][0]["rowsCapped"] is True  # 2 rows == MAX_TOP


def test_groups_for_owner_counts_and_skips_empty(monkeypatch):
    async def fake_query(entity, **kw):
        assert entity == "Case" and kw["filter"] == f"Owner/Id eq {C}"
        return [{"Group": {"Id": G, "Name": "Custom"}}, {"Group": {"Id": G, "Name": "Custom"}},
                {"Group": {"Id": "00000000-0000-0000-0000-000000000000", "Name": ""}}]

    monkeypatch.setattr(lifecycle, "query_records", fake_query)
    assert asyncio.run(lifecycle.groups_for_owner(C)) == [{"Id": G, "Name": "Custom", "count": 2}]


@pytest.fixture
def client():
    with TestClient(server.app) as c:
        yield c


def test_routes_return_400_on_bad_input(client):
    r = client.get("/api/lifecycle/groups?owner=x")
    assert r.status_code == 400 and r.json()["message"] == "Invalid owner id."
    r = client.post("/api/lifecycle", json={"groupId": "x", "since": "2025-09-01"})
    assert r.status_code == 400


def sample_result():
    rows = lifecycle.shape_rows(RAW)
    case = {"Id": C, "Number": "SR1", "Subject": "Report </script> card", "CreatedOn": "2026-01-01T00:00:00Z",
            "ClosureDate": "2026-01-04T00:00:00Z", "DateNeeded": None, "Owner": "Bo", "SubStatus": "",
            "rows": rows, "rowsCapped": False, **lifecycle.analyze_case(rows)}
    return {"cases": [case], "summary": lifecycle.summarize([case])}


def parse_sse(text):
    out = []
    for block in text.split("\n\n"):
        ev = next((l[6:].strip() for l in block.split("\n") if l.startswith("event:")), None)
        data = "".join(l[5:].strip() for l in block.split("\n") if l.startswith("data:"))
        if ev and data:
            out.append((ev, json.loads(data)))
    return out


def test_route_streams_progress_then_result(client, monkeypatch):
    monkeypatch.setattr(lifecycle, "ALLOWED_ENTITIES", [])

    async def fake_sample(gid, since, n, progress):
        assert (gid, since, n) == (G, "2025-09-01", 5)
        progress({"phase": "cases", "done": 0, "total": 5})
        progress({"phase": "rows", "done": 1, "total": 1})
        return sample_result()

    monkeypatch.setattr(server, "lifecycle_sample", fake_sample)
    r = client.post("/api/lifecycle", json={"groupId": G, "groupName": "Custom", "since": "2025-09-01", "maxCases": 5})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/event-stream")
    events = parse_sse(r.text)
    assert [e for e, _ in events] == ["progress", "progress", "result"]
    result = events[-1][1]
    assert result["summary"]["cases"] == 1 and result["runId"] in server._LC_RUNS

    # The finished pull can then be exported as Excel.
    x = client.get(f"/api/lifecycle/export?run={result['runId']}")
    assert x.status_code == 200
    assert x.headers["content-disposition"] == 'attachment; filename="Custom lifecycle.xlsx"'
    assert x.content[:2] == b"PK"  # a zip, i.e. an .xlsx


def test_export_unknown_run_is_404(client):
    r = client.get("/api/lifecycle/export?run=nope")
    assert r.status_code == 404 and "Pull lifecycle again" in r.json()["message"]


def test_xlsx_has_sheets_and_formulas():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(lifecycle_export.build_xlsx(sample_result(), "Custom", "2025-09-01")))
    assert wb.sheetnames == ["Read me", "Cases", "Lifecycle rows"]
    cases, rows = wb["Cases"], wb["Lifecycle rows"]
    assert cases["A2"].value == "SR1" and cases["I2"].value.startswith("=COUNTIF(")
    assert rows.max_row == 1 + len(RAW)
    assert rows["H2"].value == '=IF(G2="","",(G2-F2)*24)'
    assert rows.cell(rows.max_row, 7).value is None  # the Closed sentinel stays blank


def test_report_html_inlines_data_safely():
    html = lifecycle_export.build_report_html(sample_result(), "Custom", "2025-09-01")
    assert "<title>Custom Lifecycle</title>" in html
    assert "__DATA__" not in html and "Report <\\/script> card" in html
    assert html.count("</script>") == 1


def test_publish_returns_artifact_url(client, monkeypatch, tmp_path):
    monkeypatch.setattr(server, "claude_available", lambda: True)
    monkeypatch.setattr(lifecycle_export._paths, "ANALYSIS_DIR", tmp_path)
    server._LC_RUNS["r1"] = {"result": sample_result(), "group": "Custom", "since": "2025-09-01"}
    seen = {}

    def fake_run(spec, on_chunk, on_done, on_error, on_tool_use=None):
        seen["spec"] = spec
        on_done({"resultText": "Published: https://claude.ai/artifact/AbC123"})

    monkeypatch.setattr(lifecycle_export, "run_claude", fake_run)
    r = client.post("/api/lifecycle/publish", json={"runId": "r1"})
    assert r.status_code == 200 and r.json() == {"url": "https://claude.ai/artifact/AbC123"}
    assert seen["spec"].tools.set == ["Read", "Artifact"]
    # Without this the CLI started from start-app.bat has no Artifact tool.
    assert seen["spec"].env == {"CLAUDE_CODE_ENTRYPOINT": "claude-desktop"}
    assert (tmp_path / "lifecycle" / "r1" / "lifecycle-report.html").is_file()


def fake_reply(text, tools=()):
    def run(spec, on_chunk, on_done, on_error, on_tool_use=None):
        for t in tools:
            on_tool_use({"name": t, "input": {}})
        on_done({"resultText": text})
    return run


def test_publish_without_artifact_tool_says_so(client, monkeypatch, tmp_path):
    monkeypatch.setattr(server, "claude_available", lambda: True)
    monkeypatch.setattr(lifecycle_export._paths, "ANALYSIS_DIR", tmp_path)
    server._LC_RUNS["r2"] = {"result": sample_result(), "group": "Custom", "since": "2025-09-01"}
    monkeypatch.setattr(lifecycle_export, "run_claude", fake_reply("no Artifact tool is available", tools=["Read"]))
    r = client.post("/api/lifecycle/publish", json={"runId": "r2"})
    assert r.status_code == 500 and "didn't offer its Artifact tool" in r.json()["message"]


def test_publish_artifact_called_but_no_url(client, monkeypatch, tmp_path):
    monkeypatch.setattr(server, "claude_available", lambda: True)
    monkeypatch.setattr(lifecycle_export._paths, "ANALYSIS_DIR", tmp_path)
    server._LC_RUNS["r3"] = {"result": sample_result(), "group": "Custom", "since": "2025-09-01"}
    monkeypatch.setattr(lifecycle_export, "run_claude", fake_reply("sorry", tools=["Read", "Artifact"]))
    r = client.post("/api/lifecycle/publish", json={"runId": "r3"})
    assert r.status_code == 500 and "artifact link" in r.json()["message"]


def test_report_page_served_locally(client):
    server._LC_RUNS["r4"] = {"result": sample_result(), "group": "Custom", "since": "2025-09-01"}
    r = client.get("/api/lifecycle/report?run=r4")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/html")
    assert "<title>Custom Lifecycle</title>" in r.text
    assert client.get("/api/lifecycle/report?run=nope").status_code == 404
